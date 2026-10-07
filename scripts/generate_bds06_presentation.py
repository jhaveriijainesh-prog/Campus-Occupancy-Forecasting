from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PPTX_PATH = ROOT / "BDS06_Final_Project_Presentation.pptx"
IMG_DIR = ROOT / "docs" / "evidence" / "screenshots"

NAVY = RGBColor(11, 36, 59)
BLUE = RGBColor(24, 90, 157)
SKY = RGBColor(217, 234, 246)
GOLD = RGBColor(195, 140, 58)
GRAY = RGBColor(84, 96, 110)
LIGHT = RGBColor(245, 247, 249)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(23, 123, 82)
RED = RGBColor(176, 58, 46)


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.35))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.55), Inches(12.0), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = NAVY

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(12.0), Inches(0.4))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.alignment = PP_ALIGN.LEFT
        run2 = p2.runs[0]
        run2.font.size = Pt(12)
        run2.font.color.rgb = GRAY


def add_footer(slide, text="BDS-06 | Verified local capstone prototype"):
    foot = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.0), Inches(0.25))
    tf = foot.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY


def add_bullets(slide, bullets, left=0.8, top=1.8, width=5.8, height=4.7, size=19):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.runs[0]
        run.font.size = Pt(size)
        run.font.color.rgb = NAVY


def add_callout(slide, text, left, top, width, height, fill_color=SKY, text_color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = BLUE

    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = text_color


def add_image(slide, image_name, left, top, width, height, caption=None):
    path = IMG_DIR / image_name
    if not path.exists():
        raise FileNotFoundError(f"Missing screenshot: {path}")
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    if caption:
        cap = slide.shapes.add_textbox(Inches(left), Inches(top + height + 0.1), Inches(width), Inches(0.35))
        tf = cap.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = caption
        run = p.runs[0]
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "BDS-06: Campus Occupancy Forecasting", "T.Y. B.Sc. Data Science | Capstone Presentation")

# top accent
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.7), Inches(13.333), Inches(0.08))
accent.fill.solid(); accent.fill.fore_color.rgb = BLUE; accent.line.fill.background()

box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), Inches(5.2), Inches(2.6))
box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = BLUE
text = box.text_frame
text.word_wrap = True
p = text.paragraphs[0]
p.text = "Project Code: BDS-06\nAcademic Level: T.Y. B.Sc. Data Science\nScope: Verified local startup MVP\nPrototype: FastAPI + Streamlit + Docker"
p.alignment = PP_ALIGN.LEFT
for r in p.runs:
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY

# right-side callout
c1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(2.3), Inches(5.3), Inches(1.2))
c1.fill.solid(); c1.fill.fore_color.rgb = SKY; c1.line.color.rgb = BLUE
c1tf = c1.text_frame; c1p = c1tf.paragraphs[0]; c1p.text = "Verified runtime\nDocker + API + Dashboard active"; c1p.alignment = PP_ALIGN.CENTER
for r in c1p.runs: r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = NAVY

c2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(3.9), Inches(5.3), Inches(1.2))
c2.fill.solid(); c2.fill.fore_color.rgb = LIGHT; c2.line.color.rgb = BLUE
c2tf = c2.text_frame; c2p = c2tf.paragraphs[0]; c2p.text = "Verified scope\n1-room, 1-hour, point forecast"; c2p.alignment = PP_ALIGN.CENTER
for r in c2p.runs: r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = NAVY

placeholder = slide.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(8.0), Inches(1.0))
ptf = placeholder.text_frame; pp = ptf.paragraphs[0]
pp.text = "Student Name: [Insert student name]\nGuide / Institution: [Insert guide name and institution]"
pp.alignment = PP_ALIGN.LEFT
for r in pp.runs:
    r.font.size = Pt(12)
    r.font.color.rgb = GRAY

add_footer(slide, "BDS-06 | Verified local capstone prototype | 2026-10-02")

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Problem Statement", "Why campus space planning needs better visibility")
add_bullets(slide, [
    "Nominal schedules and enrollment alone do not show actual room usage.",
    "Facilities teams need better visibility into utilization, wasted capacity, and near-term demand.",
    "This project addresses the problem on a synthetic campus dataset with a verified local prototype.",
    "The scope is decision support, not deployment to a live institutional environment.",
], left=0.8, top=1.8, width=6.1, height=4.3, size=20)
callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.4), Inches(2.1), Inches(4.8), Inches(3.2))
callout.fill.solid(); callout.fill.fore_color.rgb = SKY; callout.line.color.rgb = BLUE
ctf = callout.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Key issue\nRoom occupancy is not the same as timetable availability."; ctp.alignment = PP_ALIGN.CENTER
for r in ctp.runs: r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = NAVY
add_footer(slide)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Objectives", "What the project set out to demonstrate")
add_bullets(slide, [
    "Build a reproducible campus occupancy analytics workflow.",
    "Forecast one-room occupancy for a one-hour horizon using a causal feature pipeline.",
    "Measure utilization using verified metrics: SUR, RFU, and WSH.",
    "Explore room grouping behavior through analytical clustering.",
    "Compare heuristic and optimization-based allocation on synthetic data.",
    "Deliver a secure, local dashboard and API for read-only operational review.",
], left=0.8, top=1.8, width=7.0, height=4.8, size=18)
callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(2.0), Inches(4.2), Inches(2.4))
callout.fill.solid(); callout.fill.fore_color.rgb = LIGHT; callout.line.color.rgb = BLUE
ctf = callout.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Verified scope\nRead-only local MVP"; ctp.alignment = PP_ALIGN.CENTER
for r in ctp.runs: r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = NAVY
add_footer(slide)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Proposed Solution", "Modular analytics and decision support pipeline")
# pipeline boxes
boxes = [
    ("1. Data intake", "Synthetic room, timetable, event, and occupancy data", 0.6, 1.9, 2.1, 1.2),
    ("2. Cleaning & validation", "Schema checks, timestamps, and data integrity validation", 2.85, 1.9, 2.2, 1.2),
    ("3. Feature engineering", "Calendar, lags, room context, rolling behavior, causal features", 5.25, 1.9, 2.1, 1.2),
    ("4. Forecasting", "XGBoost one-hour point forecast with chronological holdout", 7.65, 1.9, 2.1, 1.2),
    ("5. Analytics", "SUR, RFU, WSH, clustering, scenario comparisons", 0.95, 3.6, 2.5, 1.2),
    ("6. API", "FastAPI health, metrics, forecast, security, and robust errors", 3.7, 3.6, 2.7, 1.2),
    ("7. Dashboard", "Overview + Forecast Explorer for local operational review", 6.7, 3.6, 3.0, 1.2),
    ("8. Decision support", "Analytical recommendations, bounded by synthetic scope", 4.1, 5.2, 4.8, 1.0),
]
for title, body, x, y, w, h in boxes:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT; shape.line.color.rgb = BLUE
    tf = shape.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]
    p.text = f"{title}\n{body}"
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs: r.font.size = Pt(12); r.font.color.rgb = NAVY
add_footer(slide)

# Slide 5 architecture
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "System Architecture", "Verified project architecture")
add_image(slide, "10_architecture.png", 0.4, 1.8, 12.5, 5.1, caption="Figure: Verified local architecture showing the FastAPI backend, Streamlit dashboard, and data/model dependencies.")
add_footer(slide)

# Slide 6 data flow
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Data Flow and Methodology", "Synthetic data to forecast and dashboard")
add_image(slide, "11_data_flow.png", 0.45, 1.8, 12.4, 5.0, caption="Figure: Data generation, validation, model inference, and dashboard delivery pipeline.")
add_footer(slide)

# Slide 7 forecasting approach
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Forecasting Approach", "Supported local forecasting boundary")
add_bullets(slide, [
    "Model: XGBoost regressor trained on synthetic room and scheduling data.",
    "Scope: verified one-room, one-hour forecast only.",
    "Semantics: point estimate; no calibrated uncertainty intervals are claimed.",
    "Leakage control: chronological train/validation/test split and causal feature logic.",
    "Example: room B01-R101, forecast horizon 1 hour, output around 0.6201575398445129.",
    "This is a short-term operational estimate, not a broad multi-horizon system.",
], left=0.8, top=1.8, width=7.0, height=4.8, size=18)
callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.4), Inches(2.1), Inches(4.0), Inches(2.7))
callout.fill.solid(); callout.fill.fore_color.rgb = SKY; callout.line.color.rgb = BLUE
ctf = callout.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Verified forecast\nRoom B01-R101\n1-hour horizon\nPoint estimate"; ctp.alignment = PP_ALIGN.CENTER
for r in ctp.runs: r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = NAVY
add_footer(slide)

# Slide 8 analytics optimization
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Analytics and Optimization", "Local analytical capabilities and bounded comparisons")
add_bullets(slide, [
    "Analytics: SUR, RFU, and WSH are implemented and verified through the API for campus/building/room views.",
    "Clustering: seeded room behavior clustering is implemented as an analytical capability, not a live dashboard page.",
    "Optimization: greedy baseline compared with CBC MILP on synthetic assignment data.",
    "Evidence: feasible solutions were produced; the same objective value was observed in one local trial.",
    "These are analytical decision-support functions, not unsupported public operational guarantees.",
], left=0.8, top=1.8, width=7.5, height=4.8, size=18)
callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(2.3), Inches(3.6), Inches(2.2))
callout.fill.solid(); callout.fill.fore_color.rgb = LIGHT; callout.line.color.rgb = BLUE
ctf = callout.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Important\nNot a separate live dashboard page"; ctp.alignment = PP_ALIGN.CENTER
for r in ctp.runs: r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = NAVY
add_footer(slide)

# Slide 9 dashboard
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Dashboard", "Overview and Forecast Explorer as the user-facing MVP")
add_image(slide, "01_overview_full.png", 0.5, 1.8, 4.0, 2.4)
add_image(slide, "02_overview_kpis.png", 4.8, 1.8, 3.7, 2.4)
add_image(slide, "03_forecast_explorer.png", 8.9, 1.8, 3.9, 2.4)
caption = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(12.0), Inches(0.35))
ctf = caption.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Verified dashboard evidence: Overview page, KPI summary, and Forecast Explorer request flow."; ctp.alignment = PP_ALIGN.LEFT; run = ctp.runs[0]; run.font.size = Pt(9); run.font.color.rgb = GRAY
add_footer(slide)

# Slide 10 live forecast demonstration
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Live Forecast Demonstration", "Verified local forecast output")
add_image(slide, "04_valid_forecast.png", 0.5, 1.8, 5.9, 3.5)
add_image(slide, "05_forecast_result_details.png", 6.8, 1.8, 5.8, 3.5)
caption = slide.shapes.add_textbox(Inches(0.8), Inches(5.6), Inches(12.0), Inches(0.35))
ctf = caption.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Room B01-R101 | Horizon: 1 hour | Prediction: point estimate only | Verified local runtime evidence."; ctp.alignment = PP_ALIGN.LEFT; run = ctp.runs[0]; run.font.size = Pt(10); run.font.color.rgb = GRAY
add_footer(slide)

# Slide 11 security robustness
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Security and Robustness", "Verified boundaries for the local MVP")
add_image(slide, "07_api_health.png", 6.6, 1.8, 5.8, 3.8)
add_bullets(slide, [
    "API key enforcement is active for protected routes.",
    "Missing/invalid credentials return safe 401 responses.",
    "A read-only key was denied optimizer access with a 403 response.",
    "Unknown room and invalid input are handled without exposing tracebacks.",
    "Health and readiness checks confirm the service is operational.",
    "This is a controlled local prototype boundary, not a public production identity system.",
], left=0.7, top=1.8, width=5.7, height=4.6, size=17)
add_footer(slide)

# Slide 12 testing reproducibility
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Testing and Reproducibility", "Verified runtime and validation evidence")
add_image(slide, "08_docker_runtime.png", 0.5, 1.8, 5.9, 3.8)
add_image(slide, "09_test_results.png", 6.7, 1.8, 5.8, 3.8)
caption = slide.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(12.0), Inches(0.35))
ctf = caption.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Evidence: Docker local runtime running, and automated validation reported 178 passed with 1 warning in 53.36 seconds."; ctp.alignment = PP_ALIGN.LEFT; run = ctp.runs[0]; run.font.size = Pt(10); run.font.color.rgb = GRAY
add_footer(slide)

# Slide 13 results
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Results", "What the verified project actually demonstrates")
add_bullets(slide, [
    "Verified runtime: Docker Compose stack, API health, dashboard readiness, and dashboard-to-API communication all succeeded locally.",
    "Forecast capability: one-room, one-hour, point-estimate forecast works in the checked-in runtime.",
    "Dashboard functionality: Overview and Forecast Explorer operate with real API-backed data and safe error states.",
    "Security boundaries: access control and dependency checks are implemented and tested.",
    "Leakage controls: the project includes chronological validation and causal feature checks.",
    "Known limitation: the result is a synthetic, local prototype rather than a production campus deployment.",
], left=0.8, top=1.8, width=7.3, height=4.8, size=18)
callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(2.0), Inches(3.9), Inches(2.5))
callout.fill.solid(); callout.fill.fore_color.rgb = SKY; callout.line.color.rgb = BLUE
ctf = callout.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Verified result\nLocal prototype works"; ctp.alignment = PP_ALIGN.CENTER
for r in ctp.runs: r.font.size = Pt(23); r.font.bold = True; r.font.color.rgb = NAVY
add_footer(slide)

# Slide 14 limitations future work
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Limitations and Future Work", "What is not yet evidenced and what could follow")
add_bullets(slide, [
    "Synthetic data only; no real campus sensor or schedule integration is proven in this repository.",
    "Forecast scope is intentionally limited to one-hour point estimates; multi-horizon and interval forecasts are deferred.",
    "The optimizer and clustering are analytical capabilities, not separate dashboard pages or validated production workflows.",
    "No public cloud deployment or institutional approval is claimed.",
    "Future work would require real-data validation, broader forecasting, stronger calibration, and full operational security review.",
], left=0.8, top=1.8, width=7.3, height=4.8, size=18)
callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(2.2), Inches(3.9), Inches(1.8))
callout.fill.solid(); callout.fill.fore_color.rgb = LIGHT; callout.line.color.rgb = BLUE
ctf = callout.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Limitation\nEvidence remains bounded"; ctp.alignment = PP_ALIGN.CENTER
for r in ctp.runs: r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = NAVY
add_footer(slide)

# Slide 15 conclusion
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
add_title(slide, "Conclusion", "Academic and engineering value of the project")
add_bullets(slide, [
    "The project demonstrates a credible local campus occupancy analytics workflow and a working decision-support prototype.",
    "From a software engineering perspective, it shows a clean modular architecture, operational API health checks, secure boundaries, and reproducible local startup.",
    "From the data science perspective, it demonstrates causal forecasting logic, leakage-aware validation, and meaningful utilization analytics.",
    "The core value is evidence-based local planning support within a clearly bounded prototype scope.",
    "Thank You",
    "Questions & Discussion",
], left=0.8, top=1.8, width=7.4, height=4.8, size=18)
callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(2.0), Inches(3.5), Inches(2.1))
callout.fill.solid(); callout.fill.fore_color.rgb = SKY; callout.line.color.rgb = BLUE
ctf = callout.text_frame; ctp = ctf.paragraphs[0]; ctp.text = "Final message\nVerified local evidence\nAcademic prototype"; ctp.alignment = PP_ALIGN.CENTER
for r in ctp.runs: r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = NAVY
add_footer(slide, "BDS-06 | Thank You | Questions & Discussion")

# Save
prs.save(str(PPTX_PATH))
print(f"Created {PPTX_PATH}")
print(f"Slides: {len(prs.slides)}")
