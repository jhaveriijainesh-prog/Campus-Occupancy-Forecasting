from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "Final_BDS06_Project_Report.docx"
IMG_DIR = ROOT / "docs" / "evidence" / "screenshots"


def add_field(paragraph, instr_text):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr_text)
    paragraph._p.append(fld)


def add_table_of_contents(doc):
    doc.add_heading("Table of Contents", level=1)
    p = doc.add_paragraph()
    p.style = "Heading 2"
    p.add_run("Table of Contents")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    p._p.append(fld)

    doc.add_paragraph()
    doc.add_heading("List of Figures", level=1)
    p = doc.add_paragraph()
    p.style = "Heading 2"
    p.add_run("List of Figures")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "LISTOFFIGURES \\h \\z \\u")
    p._p.append(fld)

    doc.add_paragraph()
    doc.add_heading("List of Tables", level=1)
    p = doc.add_paragraph()
    p.style = "Heading 2"
    p.add_run("List of Tables")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "LISTOFTABLES \\h \\z \\u")
    p._p.append(fld)

    doc.add_page_break()


def add_image(doc, image_name, caption, width=Inches(6.8), new_page=False):
    if new_page:
        doc.add_page_break()
    image_path = IMG_DIR / image_name
    if not image_path.exists():
        raise FileNotFoundError(f"Missing evidence screenshot: {image_path}")
    doc.add_picture(str(image_path), width=width)
    doc.add_paragraph(caption)
    doc.add_paragraph()


def add_chapter_heading(doc, title, number):
    doc.add_page_break()
    doc.add_heading(f"Chapter {number}: {title}", level=1)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("Campus Occupancy Forecasting")
    title_run.bold = True
    title_run.font.size = Pt(26)

    subtitle_para = doc.add_paragraph()
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = subtitle_para.add_run("BDS-06 Final Project Report")
    subtitle.bold = True
    subtitle.font.size = Pt(18)

    doc.add_paragraph()
    doc.add_paragraph("Programme: B.Sc. Data Science")
    doc.add_paragraph("Project Code: BDS-06")
    doc.add_paragraph("Submission Type: Post-Capstone Final Report and Evidence Package")
    doc.add_paragraph("Date: 2026-10-01")
    doc.add_paragraph("Prepared for: Academic assessment and portfolio review")
    doc.add_paragraph()
    doc.add_paragraph("Scope note: This report reflects verified local runtime evidence and repository artifacts only. It does not claim public hosting or live institutional deployment.")
    doc.add_paragraph("Institutional placeholders: insert guide/supervisor name, signatures, and approval details before final submission.")

    doc.add_page_break()
    doc.add_heading("Declaration and Approval", level=1)
    doc.add_paragraph("I declare that this report is based on the verified local prototype and repository evidence captured for BDS-06. Any institutional signatures, approval marks, and supervisor details should be completed in the final submission copy.")
    doc.add_paragraph("Student/Author: [Insert student name]")
    doc.add_paragraph("Guide/Supervisor: [Insert guide name]")
    doc.add_paragraph("Signature: _________________________")
    doc.add_paragraph("Date: _________________________")

    doc.add_page_break()
    doc.add_heading("Abstract", level=1)
    abstract = (
        "BDS-06 implements a local campus occupancy analytics prototype based on synthetic room, timetable, and hourly occupancy data. "
        "The system includes data validation and cleaning, utilization metrics, a one-hour XGBoost forecast, a FastAPI service, and a Streamlit dashboard. "
        "This report documents the verified startup MVP and the supporting evidence package. "
        "The results show that the project is functionally working in a local Docker Compose environment and that the dashboard and API communicate correctly, but the scope remains intentionally limited to a read-only local prototype. "
        "The forecast is supported only for a one-hour point estimate, and no real campus deployment, stakeholder study, or calibrated uncertainty outputs are claimed."
    )
    doc.add_paragraph(abstract)
    doc.add_paragraph("Keywords: campus occupancy, forecasting, capacity optimization, synthetic data, local dashboard, FastAPI, XGBoost")

    add_table_of_contents(doc)

    add_chapter_heading(doc, "Introduction and Problem Context", 1)
    doc.add_paragraph(
        "Campus facilities planning depends on visibility into room use and short-term occupancy. The prototype addresses this by combining synthetic timetable, room, and occupancy data with a small dashboard and a one-hour forecast workflow. The objective is to help planners inspect current utilization patterns and estimate likely occupancy for a selected room without overclaiming operational deployment."
    )
    doc.add_paragraph(
        "The current workspace does not contain real campus sensor data or institutional scheduling data. The project therefore uses a generated synthetic dataset and makes a clear distinction between verified local demo behavior and unsupported production claims."
    )
    doc.add_paragraph(
        "The design scope remains intentionally narrow: a read-only dashboard with Overview and Forecast Explorer pages, health checks, and local API-backed metrics and forecast requests. The repository also contains analysis and optimization capabilities that remain outside the current UI MVP."
    )

    add_chapter_heading(doc, "Requirements, Stakeholders, and Verified Scope", 2)
    doc.add_paragraph("The project serves facilities and space-planning users, technical operators, and academic evaluators. The principal user journeys are reading readiness, inspecting KPI rollups, selecting a room, and generating a one-hour occupancy forecast. No formal user study is included in the repository evidence.")
    doc.add_paragraph("The verified current scope is: health/readiness and API access, Dashboard Overview and Forecast Explorer, synthetic processed artifacts, one-hour XGBoost forecast, and a local Docker Compose runtime. The following items are not current verified capability: public deployment, multi-horizon forecast, calibrated prediction intervals, real campus integration, and optimizer/simulation dashboard pages.")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Area"
    header[1].text = "Verified status"
    header[2].text = "Evidence"
    header[3].text = "Comment"
    rows = [
        ["Health and API", "Verified locally", "/api/v1/health, /api/v1/health/ready", "Service status and readiness passed"],
        ["Overview dashboard", "Verified locally", "01_overview_full.png and 02_overview_kpis.png", "Readiness + KPI summary shown"],
        ["Forecast explorer", "Verified locally", "03_forecast_explorer.png and 04_valid_forecast.png", "Room forecast request was successful"],
        ["Forecast scope", "One-hour point forecast only", "model/info and forecast artifact", "No multi-horizon or calibrated uncertainty claim"],
        ["Optimization", "Code + offline evidence only", "solver and heuristic modules", "Not promoted as dashboard page"],
        ["Deployment", "Local Docker only", "docker compose ps and runtime checks", "No public hosting claim"],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value

    doc.add_paragraph("The runtime and repository evidence support the startup MVP as a verified local product boundary, not a production campus service.")

    add_chapter_heading(doc, "System Architecture and Data Flow", 3)
    doc.add_paragraph("The implementation is a modular Python monolith composed of data, feature, forecasting, analytics, optimization, API, and dashboard layers. FastAPI serves the runtime API, while Streamlit provides the dashboard. The two services communicate over Docker Compose internal DNS in the local environment.")
    doc.add_paragraph("This architecture supports a read-only operational workflow: generate, validate, and process synthetic occupancy data; load the XGBoost model; serve metrics and forecast endpoints; and expose the dashboard to a user without exposing internal file paths or unsupported operations.")
    add_image(doc, "10_architecture.png", "Figure 3.1: Verified architecture diagram showing the local FastAPI backend, Streamlit dashboard, and the data/model dependencies of the startup MVP.", Inches(6.6))
    add_image(doc, "11_data_flow.png", "Figure 3.2: Verified data flow from synthetic generation to cleaning, model inference, and API/dashboard delivery.", Inches(6.6))

    add_chapter_heading(doc, "Implementation Evidence and Live Runtime", 4)
    doc.add_paragraph("The core evidence package was captured from the live local runtime. The screenshots below represent the actual system state and confirm the behavior of the dashboard, API, and end-to-end runtime checks. Accordingly, the report omits any unsupported evidence files and includes only screenshots with direct project provenance.")
    image_captions = [
        ("01_overview_full.png", "Figure 4.1: Overview page showing the main dashboard shell and KPI context."),
        ("02_overview_kpis.png", "Figure 4.2: KPI summary area with readiness and utilization indicators used in the local demo."),
        ("03_forecast_explorer.png", "Figure 4.3: Forecast Explorer page before request generation."),
        ("04_valid_forecast.png", "Figure 4.4: Valid room forecast returned from the live API-backed workflow."),
        ("05_forecast_result_details.png", "Figure 4.5: Forecast result details including room, horizon, and model metadata."),
        ("07_api_health.png", "Figure 4.6: API health endpoint response from the live local service."),
        ("08_docker_runtime.png", "Figure 4.7: Docker runtime status confirming the Compose stack is running."),
        ("09_test_results.png", "Figure 4.8: Pytest execution summary showing the verified local validation pass.")
    ]
    for image_name, caption in image_captions:
        add_image(doc, image_name, caption, Inches(6.7))

    add_chapter_heading(doc, "Forecasting, Evaluation, and Metrics", 5)
    doc.add_paragraph("The forecasting component is based on an XGBoost regressor trained on the processed synthetic dataset. The served artifact supports a one-hour point forecast and does not claim calibrated uncertainty intervals or multi-horizon prediction. The evaluation design uses a chronological holdout test split to keep the assessment close to deployment conditions for a short-horizon forecast.")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Model", "MAE", "RMSE", "R2", "WAPE", "sMAPE"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    rows = [
        ["Historical seasonal profile", "2.5975", "10.1839", "0.2335", "102.72%", "163.6674%"],
        ["XGBoost", "1.2775", "4.3961", "0.8572", "50.52%", "155.5725%"],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph("These figures describe the synthetic evaluation sample and should not be interpreted as real-world campus performance. They establish that the forecast model is effective within the bounds of the checked-in data and the supported one-hour horizon.")
    doc.add_paragraph("The utilization metrics in the API include SUR, RFU, and WSH. The backend validates these formulas against the synthetic dataset and returns consistent values for campus, building, and room scope requests. The API contract remains read-only and does not imply a production-grade campus optimization engine.")

    add_chapter_heading(doc, "Limitations, Risks, and Operational Guidance", 6)
    doc.add_paragraph("The report intentionally keeps the project honest about risk. Real campus data, live integration, and production hosting are not part of the verified workspace state. The local demo therefore demonstrates engineering readiness and evidence quality, not institutional operational readiness.")
    doc.add_paragraph("The main risks are: lack of live sensor data, unsupported public deployment status, the one-hour forecast ceiling, no calibrated uncertainty bands, and no formal user acceptance testing. The recommended operational practice is to keep the system local, use the documented API and dashboard, and keep all model claims aligned with the checked-in synthetic artifacts.")
    doc.add_paragraph("Before any live deployment or institutional rollout, the team would need a controlled data governance review, procurement/approval process, stakeholder evaluation, and a revalidated model package built on a real campus dataset rather than the synthetic generator used in the project repository.")

    add_chapter_heading(doc, "Conclusion and Reproducibility", 7)
    doc.add_paragraph("The BDS-06 project has delivered a verified local campus occupancy analytic prototype with a functional API, a dashboard, and supporting evaluation evidence. The work demonstrates that the core workflow is operational for the local startup MVP: monitoring readiness, calculating utilization metrics, generating one-hour forecasts, and validating runtime behavior through Docker and automated tests.")
    doc.add_paragraph("The final evidence package confirms that the system is ready for local review and portfolio demonstration, while maintaining the correct caution: it is not a production service, not a public deployment, and not a claim of live campus deployment or real-world operational savings.")
    doc.add_paragraph("Reproducibility is supported through the repository structure, Docker Compose configuration, local scripts, evaluation artifacts, and the captured runtime evidence. In the final institutional copy, the project guide should insert the required signatures and approval fields before archiving or submission.")

    doc.add_page_break()
    doc.add_heading("Appendix A: evidence traceability", level=1)
    doc.add_paragraph("Verified runtime evidence and screenshot files are stored under the docs/evidence/screenshots folder and are the authoritative source for the final report. The project explicitly avoids unsupported screenshots and claims, including any image outside the verified local evidence set.")
    doc.add_paragraph("The final report was generated from the repository evidence and the live local system state. The screenshots used in the document are the same files used in the evidence package: 01_overview_full.png, 02_overview_kpis.png, 03_forecast_explorer.png, 04_valid_forecast.png, 05_forecast_result_details.png, 07_api_health.png, 08_docker_runtime.png, 09_test_results.png, 10_architecture.png, and 11_data_flow.png.")
    doc.add_paragraph("The Word export captures the verified implementation boundary rather than introducing any new or inflated capabilities.")

    return doc


def main():
    doc = build_document()
    doc.save(str(OUTPUT_PATH))
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
