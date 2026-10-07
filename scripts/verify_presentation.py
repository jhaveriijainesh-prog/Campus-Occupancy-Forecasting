import zipfile
from pathlib import Path

from pptx import Presentation

pptx_path = Path(__file__).resolve().parents[1] / "BDS06_Final_Project_Presentation.pptx"
prs = Presentation(str(pptx_path))
print("slides=", len(prs.slides))
for index, slide in enumerate(prs.slides, 1):
    title = slide.shapes.title.text if slide.shapes.title else "(no title)"
    print(index, title)

with zipfile.ZipFile(pptx_path) as zf:
    media = sorted(name for name in zf.namelist() if name.startswith("ppt/media/"))
    print("media_files=", len(media))
    print("media_names=", media)
